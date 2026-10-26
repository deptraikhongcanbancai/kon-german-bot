import os
import re
from io import BytesIO
from pathlib import Path
import requests
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from config import KON_EXPORT_GIFS

GIFS_BY_SEVERITY = {
    "strong": [
        "https://media.giphy.com/media/10JhviFuU2gWD6/giphy.gif",
        "https://media.giphy.com/media/l3q2K5jinAlChoCLS/giphy.gif",
        "https://media.giphy.com/media/3o7TKWineS040erhjq/giphy.gif",
        "https://media.giphy.com/media/26ufdipQqU2lhNA4g/giphy.gif",
    ],
    "medium": [
        "https://media.giphy.com/media/B37cYPCruqOU8/giphy.gif",
        "https://media.giphy.com/media/l4JyX3V0yydvPHNBe/giphy.gif",
        "https://media.giphy.com/media/3oEjI6SIIHBdRxXI40/giphy.gif",
        "https://media.giphy.com/media/9Y5BbDSkSTiY8/giphy.gif",
    ],
    "light": [
        "https://media.giphy.com/media/Q7ozWVYCR0nyW2rvPW/giphy.gif",
        "https://media.giphy.com/media/3o6ZtaO9BZHcOjmErm/giphy.gif",
        "https://media.giphy.com/media/xUPGcguWZHRC2HyBRS/giphy.gif",
        "https://media.giphy.com/media/L95W4wv8nnb9K/giphy.gif",
    ],
}

GIFS_BY_ERROR_TYPE = {
    "article_case": [
        "https://media.giphy.com/media/WRQBXSCnEFJIuxktnw/giphy.gif",
        "https://media.giphy.com/media/3o6Zt481isNVuQI1l6/giphy.gif",
        "https://media.giphy.com/media/xT0xeJpnrWC4XWblEk/giphy.gif",
        "https://media.giphy.com/media/l0IylOPCNkiqOgMyA/giphy.gif",
        "https://media.giphy.com/media/13GIgrGdslD9oQ/giphy.gif",
        "https://media.giphy.com/media/3owzW5c1tPq63MPmWk/giphy.gif",
    ],
    "verb_conjugation": [
        "https://media.giphy.com/media/26ufdipQqU2lhNA4g/giphy.gif",
        "https://media.giphy.com/media/3oEjI789af0AVurF60/giphy.gif",
        "https://media.giphy.com/media/l3q2K5jinAlChoCLS/giphy.gif",
        "https://media.giphy.com/media/3o7TKWineS040erhjq/giphy.gif",
        "https://media.giphy.com/media/xT9IgIc0lryrxvqVGM/giphy.gif",
        "https://media.giphy.com/media/1L5YuA6wpKkNO/giphy.gif",
    ],
    "word_order": [
        "https://media.giphy.com/media/3o6Zt6ML6BklcajjsA/giphy.gif",
        "https://media.giphy.com/media/3o7btPCcdNniyf0ArS/giphy.gif",
        "https://media.giphy.com/media/l0HlvtIPzPdt2usKs/giphy.gif",
        "https://media.giphy.com/media/9Y5BbDSkSTiY8/giphy.gif",
        "https://media.giphy.com/media/3ohzdYJK1wAdPWVk88/giphy.gif",
        "https://media.giphy.com/media/4JVTF9zR9BicshFAb7/giphy.gif",
    ],
    "preposition": [
        "https://media.giphy.com/media/xUPGcguWZHRC2HyBRS/giphy.gif",
        "https://media.giphy.com/media/l4JyX3V0yydvPHNBe/giphy.gif",
        "https://media.giphy.com/media/3ohhwF34cGDoFFhRfy/giphy.gif",
        "https://media.giphy.com/media/3o6ZtaO9BZHcOjmErm/giphy.gif",
        "https://media.giphy.com/media/1zSz5MVw4zKg0/giphy.gif",
        "https://media.giphy.com/media/26gsjCZpPolPr3sBy/giphy.gif",
    ],
    "vocabulary_style": [
        "https://media.giphy.com/media/B37cYPCruqOU8/giphy.gif",
        "https://media.giphy.com/media/L95W4wv8nnb9K/giphy.gif",
        "https://media.giphy.com/media/Q7ozWVYCR0nyW2rvPW/giphy.gif",
        "https://media.giphy.com/media/10JhviFuU2gWD6/giphy.gif",
        "https://media.giphy.com/media/3o7aD4kZn5k0SEvPmo/giphy.gif",
        "https://media.giphy.com/media/3o7abKhOpu0NwenH3O/giphy.gif",
    ],
    "spelling": [
        "https://media.giphy.com/media/3o7aD2saalBwwftBIY/giphy.gif",
        "https://media.giphy.com/media/3o7btPCcdNniyf0ArS/giphy.gif",
        "https://media.giphy.com/media/3o6Zt481isNVuQI1l6/giphy.gif",
        "https://media.giphy.com/media/26ufdipQqU2lhNA4g/giphy.gif",
        "https://media.giphy.com/media/3oEduOnl5IHM5NRodO/giphy.gif",
        "https://media.giphy.com/media/3o6nV9IUiNn1LCjSak/giphy.gif",
    ],
    "structure": [
        "https://media.giphy.com/media/26u4cqiYI30juCOGY/giphy.gif",
        "https://media.giphy.com/media/111ebonMs90YLu/giphy.gif",
        "https://media.giphy.com/media/l0HlvtIPzPdt2usKs/giphy.gif",
        "https://media.giphy.com/media/3oEjI6SIIHBdRxXI40/giphy.gif",
        "https://media.giphy.com/media/3oKIPEqDGUULpEU0aQ/giphy.gif",
        "https://media.giphy.com/media/3o7TKTDn976rzVgky4/giphy.gif",
    ],
}

ERROR_HINTS = {
    "strong": [
        "sai hoan toan",
        "sai nghiem trong",
        "khong dung",
        "sai cau truc",
        "sai chia dong tu",
        "sai word order",
        "sai thu tu tu",
        "khong tu nhien",
    ],
    "medium": [
        "can sua",
        "nen sua",
        "chua tu nhien",
        "chua dung",
        "sai gioi tu",
        "sai giong",
        "sai mao tu",
        "sai cach",
        "sai giong tu",
    ],
}

ERROR_TYPE_HINTS = {
    "article_case": [
        "artikel", "article", "mao tu", "mạo từ", "kasus", "case",
        "akkusativ", "dativ", "nominativ", "genitiv", "giong", "giống",
        "der ", "die ", "das ", "den ", "dem ",
    ],
    "verb_conjugation": [
        "chia dong tu", "chia động từ", "verb", "konjug", "conjug",
        "partizip", "perfekt", "prateritum", "präteritum", "modalverb", "hilfsverb",
    ],
    "word_order": [
        "word order", "thu tu tu", "thứ tự từ", "vi tri dong tu", "vị trí động từ",
        "satzstellung", "nebensatz", "hauptsatz", "verb am ende",
    ],
    "preposition": [
        "gioi tu", "giới từ", "preposition", "praposition", "präposition",
        "wechselpraposition", "wechselpräposition",
    ],
    "vocabulary_style": [
        "tu vung", "từ vựng", "wortschatz", "word choice", "chon tu", "chọn từ",
        "khong tu nhien", "không tự nhiên", "style", "stil",
    ],
    "spelling": [
        "rechtschreibung", "spelling", "chinh ta", "chính tả",
        "umlaut", "ß", "viet hoa", "viết hoa", "gross", "groß",
    ],
    "structure": [
        "cau truc", "cấu trúc", "struktur", "sentence", "satz",
        "mệnh đề", "menh de", "cohesion", "koharenz", "kohärenz",
    ],
}

CRITERION_NAMES = [
    "aufgabenerfullung",
    "aufgabenerfüllung",
    "aussprache",
    "intonation",
    "flussigkeit",
    "flüssigkeit",
    "koharenz",
    "kohärenz",
    "wortschatz",
    "struktur",
]

LEVEL_GIFS = {
    "excellent": [
        "https://media.giphy.com/media/111ebonMs90YLu/giphy.gif",
        "https://media.giphy.com/media/ely3apij36BJhoZ234/giphy.gif",
        "https://media.giphy.com/media/26u4cqiYI30juCOGY/giphy.gif",
        "https://media.giphy.com/media/l0IykG0AM7911MrCM/giphy.gif",
    ],
    "good": [
        "https://media.giphy.com/media/26u4cqiYI30juCOGY/giphy.gif",
        "https://media.giphy.com/media/3o7abKhOpu0NwenH3O/giphy.gif",
        "https://media.giphy.com/media/5GoVLqeAOo6PK/giphy.gif",
        "https://media.giphy.com/media/3o7aD4kZn5k0SEvPmo/giphy.gif",
    ],
    "ok": [
        "https://media.giphy.com/media/l0HlvtIPzPdt2usKs/giphy.gif",
        "https://media.giphy.com/media/3o6ZtaO9BZHcOjmErm/giphy.gif",
        "https://media.giphy.com/media/13GIgrGdslD9oQ/giphy.gif",
        "https://media.giphy.com/media/xT0xeJpnrWC4XWblEk/giphy.gif",
    ],
    "weak": [
        "https://media.giphy.com/media/3o7TKWineS040erhjq/giphy.gif",
        "https://media.giphy.com/media/WRQBXSCnEFJIuxktnw/giphy.gif",
        "https://media.giphy.com/media/l3q2K5jinAlChoCLS/giphy.gif",
        "https://media.giphy.com/media/26ufdipQqU2lhNA4g/giphy.gif",
    ],
}

VI_LINE_REWRITES = {
    "ban goc (transcription)": "Bản gốc (Transcription)",
    "bang diem": "Bảng Điểm",
    "phan tich phat am & ngu dieu (aussprache-coaching)": "Phân tích Phát âm & Ngữ điệu (Aussprache-Coaching)",
    "nhan xet & loi khuyen": "Nhận xét & Lời khuyên",
    "loi can sua / sua chi tiet": "Lỗi cần sửa / Sửa chi tiết",
    "ban hoan thien (polished version)": "Bản hoàn thiện (Polished Version)",
    "thu thach nho (mini challenge)": "Thử thách nhỏ (Mini Challenge)",
    "thu thach luyen noi (speaking challenge)": "Thử thách Luyện nói (Speaking Challenge)",
    "nguyen ban:": "Nguyên bản:",
    "sua lai:": "Sửa lại:",
    "ly do:": "Lý do:",
}


def read_docx_text(path: Path, include_headings: bool = True, max_chars: int = 15000) -> str:
    """Đọc nội dung văn bản từ file .docx người dùng gửi lên."""
    document = Document(str(path))
    chunks = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = ""
        if paragraph.style is not None and paragraph.style.name is not None:
            style_name = paragraph.style.name.lower()
        if include_headings and style_name.startswith("heading"):
            chunks.append(f"\n# {text}\n")
        else:
            chunks.append(text)

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                chunks.append(row_text)

    full_text = "\n".join(chunks).strip()
    if max_chars > 0 and len(full_text) > max_chars:
        full_text = full_text[:max_chars].rstrip() + "\n\n[ĐÃ CẮT BỚT]"
    return full_text


def set_document_defaults(doc):
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"]:
        try:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            rpr = style._element.get_or_add_rPr()
            rfonts = rpr.get_or_add_rFonts()
            rfonts.set(qn("w:ascii"), "Times New Roman")
            rfonts.set(qn("w:hAnsi"), "Times New Roman")
            rfonts.set(qn("w:eastAsia"), "Times New Roman")
            rfonts.set(qn("w:cs"), "Times New Roman")
        except Exception:
            pass


def apply_font_to_run(run):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")


def apply_font_to_paragraph(paragraph):
    for run in paragraph.runs:
        apply_font_to_run(run)


def normalize_text(text: str) -> str:
    return (
        text.lower()
        .replace("ä", "a")
        .replace("ö", "o")
        .replace("ü", "u")
        .replace("ß", "ss")
    )


def repair_mojibake_text(text: str) -> str:
    suspicious_markers = ("Ã", "Â", "Ä", "Å", "Æ", "â", "ðŸ")
    if not any(marker in text for marker in suspicious_markers):
        return text
    for source_encoding in ("cp1252", "latin-1"):
        try:
            repaired = text.encode(source_encoding).decode("utf-8")
            if repaired != text:
                return repaired
        except (UnicodeEncodeError, UnicodeDecodeError):
            continue
    return text


def rewrite_common_vi_line(line: str) -> str:
    stripped = line.strip()
    prefix = ""
    body = stripped

    if stripped.startswith("#"):
        level = len(stripped) - len(stripped.lstrip("#"))
        prefix = f"{'#' * level} "
        body = stripped[level:].strip()
    elif stripped.startswith("- "):
        prefix = "- "
        body = stripped[2:].strip()
    elif stripped.startswith("* "):
        prefix = "* "
        body = stripped[2:].strip()
    elif re.match(r"^\d+\.\s", stripped):
        match = re.match(r"^(\d+\.\s)(.*)$", stripped)
        prefix = match.group(1)
        body = match.group(2).strip()

    normalized = normalize_text(body).strip()
    if normalized in VI_LINE_REWRITES:
        return prefix + VI_LINE_REWRITES[normalized]

    if normalized.startswith("bang diem goethe"):
        return prefix + body.replace("Bang Diem", "Bảng Điểm", 1)
    if normalized.startswith("ban hoan thien"):
        return prefix + body.replace("Ban hoan thien", "Bản hoàn thiện", 1)
    if normalized.startswith("ban goc"):
        return prefix + body.replace("Ban goc", "Bản gốc", 1)
    if normalized.startswith("thu thach nho"):
        return prefix + body.replace("Thu thach nho", "Thử thách nhỏ", 1)

    for key, replacement in VI_LINE_REWRITES.items():
        if key.endswith(":") and normalized.startswith(key):
            suffix = body[len(body.split(":", 1)[0]) + 1 :]
            return prefix + replacement + suffix

    return line


def detect_error_type(text: str):
    normalized = normalize_text(text)
    for error_type, hints in ERROR_TYPE_HINTS.items():
        if any(normalize_text(hint) in normalized for hint in hints):
            return error_type
    return None


def detect_gif_severity(text: str) -> str:
    normalized = text.lower()
    for hint in ERROR_HINTS["strong"]:
        if hint in normalized:
            return "strong"
    for hint in ERROR_HINTS["medium"]:
        if hint in normalized:
            return "medium"
    return "light"


def pick_unique_from_options(line: str, options: list, used_gifs: set) -> str:
    start_idx = sum(ord(ch) for ch in line) % len(options)
    for offset in range(len(options)):
        candidate = options[(start_idx + offset) % len(options)]
        if candidate not in used_gifs:
            used_gifs.add(candidate)
            return candidate
    candidate = options[start_idx]
    used_gifs.add(candidate)
    return candidate


def pick_gif_url(line: str, severity: str, used_gifs: set) -> str:
    error_type = detect_error_type(line)
    if error_type:
        options = GIFS_BY_ERROR_TYPE[error_type] + GIFS_BY_SEVERITY[severity]
    else:
        options = GIFS_BY_SEVERITY[severity]
    return pick_unique_from_options(line, options, used_gifs)


def is_criterion_line(line: str) -> bool:
    normalized = normalize_text(line)
    return any(key in normalized for key in CRITERION_NAMES)


def criterion_level_gif(line: str, used_gifs: set) -> str:
    score_match = re.search(r"(\d+(?:[.,]\d+)?)\s*/\s*(\d+)", line)
    if not score_match:
        return pick_unique_from_options(line, LEVEL_GIFS["ok"], used_gifs)
    score = float(score_match.group(1).replace(",", "."))
    maximum = float(score_match.group(2))
    ratio = score / maximum if maximum else 0
    if ratio >= 0.85:
        return pick_unique_from_options(line, LEVEL_GIFS["excellent"], used_gifs)
    if ratio >= 0.65:
        return pick_unique_from_options(line, LEVEL_GIFS["good"], used_gifs)
    if ratio >= 0.45:
        return pick_unique_from_options(line, LEVEL_GIFS["ok"], used_gifs)
    return pick_unique_from_options(line, LEVEL_GIFS["weak"], used_gifs)


import hashlib
from config import DATA_DIR

GIF_CACHE_DIR = DATA_DIR / "gif_cache"
GIF_CACHE_DIR.mkdir(parents=True, exist_ok=True)
_GIPHY_UNREACHABLE = False


def add_picture_from_url(doc, img_url: str) -> bool:
    global _GIPHY_UNREACHABLE
    if not KON_EXPORT_GIFS:
        return False

    # Dùng bản tối ưu dung lượng /200w_s.gif (~30KB/ảnh) của Giphy để file Word nhẹ (~350KB thay vì 28MB), gửi Telegram siêu tốc
    if "giphy.com/media/" in img_url and img_url.endswith("/giphy.gif"):
        img_url = img_url[:-10] + "/200w_s.gif"

    url_hash = hashlib.md5(img_url.encode("utf-8")).hexdigest()
    cache_file = GIF_CACHE_DIR / f"{url_hash}.gif"

    if cache_file.exists() and cache_file.stat().st_size > 0:
        try:
            doc.add_picture(BytesIO(cache_file.read_bytes()), width=Cm(7.5))
            return True
        except Exception:
            pass

    if _GIPHY_UNREACHABLE:
        return False

    try:
        resp = requests.get(img_url, timeout=(2.0, 3.0))
        if resp.status_code == 200 and resp.content:
            cache_file.write_bytes(resp.content)
            doc.add_picture(BytesIO(resp.content), width=Cm(7.5))
            return True
    except Exception:
        _GIPHY_UNREACHABLE = True
    return False


def format_p(paragraph, text: str):
    text = text.replace("`", "")
    parts = re.split(r"(\*\*.*?\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        apply_font_to_run(run)


def looks_like_detail_heading(line: str) -> bool:
    lowered = normalize_text(line)
    raw = line.lower()
    return (
        "loi can sua" in lowered
        or "sua chi tiet" in lowered
        or "loi sai chinh" in lowered
        or "sửa lỗi" in raw
        or "sửa chi tiết" in raw
        or "lỗi cần sửa" in raw
    )


def looks_like_polished_heading(line: str) -> bool:
    raw = line.lower()
    lowered = normalize_text(line)
    return line.startswith("#") and (
        "ban hoan thien" in lowered
        or "bản hoàn thiện" in raw
        or "polished version" in raw
        or "bai sua tham khao" in lowered
        or "thu thach nho" in lowered
    )


def is_error_block_start(line: str) -> bool:
    if line.startswith("#"):
        return False
    lowered = normalize_text(line)
    raw = line.lower()
    return (
        bool(re.match(r"^\d+\.\s", line))
        or "nguyen ban" in lowered
        or "nguyên bản" in raw
        or "cau goc" in lowered
        or "câu gốc" in raw
    )


def markdown_to_docx(markdown_text: str, out_path: Path, title: str = "Kết Quả Chấm Bài & Giáo Án - Kon 🤓") -> Path:
    """Chuyển đổi nội dung Markdown sang file Word (.docx) kèm định dạng chuẩn và GIF reaction."""
    used_gifs = set()
    text = repair_mojibake_text(markdown_text)

    doc = Document()
    set_document_defaults(doc)
    heading = doc.add_heading(title, 1)
    apply_font_to_paragraph(heading)

    img_pattern = re.compile(r"!\[.*?\]\((https?://[^\)]+)\)")
    in_detail_section = False
    pending_detail_gif_line = None

    def add_reaction_gif_block(target_line: str):
        if not KON_EXPORT_GIFS:
            return
        severity = detect_gif_severity(target_line)
        gif_url = pick_gif_url(target_line, severity, used_gifs)
        doc.add_paragraph()
        add_picture_from_url(doc, gif_url)
        doc.add_paragraph()

    def add_line(line_str: str):
        line_str = repair_mojibake_text(rewrite_common_vi_line(line_str))
        if KON_EXPORT_GIFS and is_criterion_line(line_str):
            p = doc.add_paragraph()
            format_p(p, line_str)
            gif_url = criterion_level_gif(line_str, used_gifs)
            add_picture_from_url(doc, gif_url)
            return

        if line_str.startswith("#"):
            level = len(line_str) - len(line_str.lstrip("#"))
            p = doc.add_heading(line_str.lstrip("#").strip(), min(level, 3))
            apply_font_to_paragraph(p)
        elif line_str.startswith("- ") or line_str.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            format_p(p, line_str[2:].strip())
        elif re.match(r"^\d+\.\s", line_str):
            p = doc.add_paragraph(style="List Number")
            format_p(p, line_str.split(".", 1)[1].strip())
        else:
            p = doc.add_paragraph()
            format_p(p, line_str)

    for raw_line in text.split("\n"):
        line = repair_mojibake_text(rewrite_common_vi_line(raw_line.strip()))
        if not line:
            continue

        if looks_like_detail_heading(line):
            in_detail_section = True
        elif looks_like_polished_heading(line):
            if pending_detail_gif_line:
                add_reaction_gif_block(pending_detail_gif_line)
                pending_detail_gif_line = None
            in_detail_section = False

        if in_detail_section and is_error_block_start(line):
            if pending_detail_gif_line:
                add_reaction_gif_block(pending_detail_gif_line)
            pending_detail_gif_line = line

        img_match = img_pattern.search(line)
        if img_match:
            img_url = img_match.group(1)
            clean_line = img_pattern.sub("", line).strip()
            if clean_line:
                add_line(clean_line)
            add_picture_from_url(doc, img_url)
            continue

        add_line(line)

    if pending_detail_gif_line:
        add_reaction_gif_block(pending_detail_gif_line)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
